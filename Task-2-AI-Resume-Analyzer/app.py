import os
import re
import streamlit as st

from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document
from langchain_google_genai import ChatGoogleGenerativeAI


# ============================================================
# CONFIG
# ============================================================

load_dotenv()

st.set_page_config(
    page_title="AI Resume Analyzer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="collapsed",
)


# ============================================================
# CUSTOM CSS
# ============================================================

st.markdown(
    """
    <style>

    /* ---------- MAIN PAGE ---------- */

    .stApp {
        background:
            radial-gradient(
                circle at top left,
                rgba(99, 63, 180, 0.20),
                transparent 32%
            ),
            radial-gradient(
                circle at top right,
                rgba(0, 150, 170, 0.14),
                transparent 30%
            ),
            #070a0f;
        color: #f5f5f5;
    }

    .main .block-container {
        max-width: 1100px;
        padding-top: 35px;
        padding-bottom: 30px;
    }


    /* ---------- HERO ---------- */

    .hero-card {
        background: rgba(25, 25, 39, 0.92);
        border: 1px solid rgba(150, 130, 190, 0.30);
        border-radius: 22px;
        padding: 34px 38px;
        margin-bottom: 38px;
        box-shadow: 0 15px 40px rgba(0, 0, 0, 0.25);
    }

    .hero-title {
        font-size: 42px;
        font-weight: 800;
        margin: 0 0 12px 0;
        line-height: 1.15;
        background: linear-gradient(
            90deg,
            #c44cff,
            #ff3ccf,
            #8c7cff
        );
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
    }

    .hero-subtitle {
        color: #cfd0df;
        font-size: 17px;
        line-height: 1.7;
        margin: 0;
    }


    /* ---------- SECTION TITLE ---------- */

    .section-title {
        font-size: 27px;
        font-weight: 800;
        color: #ffffff;
        margin: 0 0 24px 0;
    }


    /* ---------- LABELS ---------- */

    .custom-label {
        color: #ffffff;
        font-size: 17px;
        font-weight: 700;
        margin-bottom: 8px;
    }

    .custom-help {
        color: #aeb1c4;
        font-size: 13px;
        margin-bottom: 8px;
    }


    /* ---------- INPUT CONTAINERS ---------- */

    [data-testid="stFileUploader"] {
        background: #f1f3f7;
        border-radius: 10px;
        padding: 6px;
    }

    [data-testid="stFileUploader"] section {
        border: none !important;
        background: transparent !important;
    }

    [data-testid="stFileUploaderDropzone"] {
        background: transparent !important;
        border: none !important;
    }

    [data-testid="stFileUploaderDropzoneInstructions"] {
        color: #6e7485 !important;
    }

    [data-testid="stFileUploaderDropzoneInstructions"] span {
        color: #6e7485 !important;
    }

    [data-testid="stFileUploader"] button {
        background: white !important;
        color: #4d5870 !important;
        border: 1px solid #cdd2dc !important;
        border-radius: 8px !important;
    }


    /* ---------- TEXT AREA ---------- */

    textarea {
        background: #f1f3f7 !important;
        color: #687188 !important;
        border: none !important;
        border-radius: 10px !important;
        font-size: 14px !important;
    }

    textarea::placeholder {
        color: #8b92a5 !important;
        opacity: 1 !important;
    }


    /* ---------- BUTTON ---------- */

    .stButton > button {
        background: linear-gradient(
            90deg,
            #8b45ff,
            #f02fc4
        ) !important;

        color: white !important;
        border: none !important;
        border-radius: 10px !important;

        padding: 12px 25px !important;

        font-size: 15px !important;
        font-weight: 700 !important;

        box-shadow:
            0 8px 25px rgba(213, 43, 210, 0.28);

        transition: all 0.2s ease;
    }

    .stButton > button:hover {
        transform: translateY(-2px);
        box-shadow:
            0 12px 30px rgba(213, 43, 210, 0.40);
    }


    /* ---------- RESULT CARDS ---------- */

    .result-card {
        background: rgba(20, 22, 31, 0.95);
        border: 1px solid #292d3b;
        border-radius: 16px;
        padding: 24px;
        margin-top: 20px;
    }

    .result-title {
        color: #ffffff;
        font-size: 21px;
        font-weight: 800;
        margin-bottom: 12px;
    }

    .score {
        font-size: 38px;
        font-weight: 800;
        background: linear-gradient(
            90deg,
            #b64cff,
            #ff3ccf
        );
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
    }

    [data-testid="stHeader"] {
    display: none !important;
}

[data-testid="stToolbar"] {
    display: none !important;
}

[data-testid="stAppViewContainer"] > .main {
    padding-top: 0 !important;
}

.block-container {
    padding-top: 2rem !important;
}


    /* ---------- FOOTER ---------- */

    .footer {
        margin-top: 70px;
        padding-top: 20px;
        border-top: 1px solid #292d38;
        text-align: center;
        color: #777d92;
        font-size: 13px;
        line-height: 1.8;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# FUNCTIONS
# ============================================================

def extract_pdf_text(uploaded_file):
    """Extract text from PDF."""
    try:
        reader = PdfReader(uploaded_file)

        pages = []

        for page in reader.pages:
            text = page.extract_text()

            if text:
                pages.append(text)

        return "\n".join(pages).strip()

    except Exception as e:
        raise Exception(f"Could not read PDF: {e}")


def extract_docx_text(uploaded_file):
    """Extract text from DOCX."""
    try:
        document = Document(uploaded_file)

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        # Also extract tables
        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        row_text.append(cell_text)

                if row_text:
                    paragraphs.append(" | ".join(row_text))

        return "\n".join(paragraphs).strip()

    except Exception as e:
        raise Exception(f"Could not read DOCX: {e}")


def extract_resume_text(uploaded_file):
    """Extract text based on file type."""

    if uploaded_file is None:
        return ""

    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        return extract_pdf_text(uploaded_file)

    if filename.endswith(".docx"):
        return extract_docx_text(uploaded_file)

    raise Exception("Unsupported file format. Please upload PDF or DOCX.")


def clean_response(text):
    """Clean Gemini response safely whether content is str or list."""

    if not text:
        return ""

    # Gemini/LangChain may return content as a list
    if isinstance(text, list):
        parts = []

        for item in text:
            if isinstance(item, str):
                parts.append(item)

            elif isinstance(item, dict):
                if "text" in item:
                    parts.append(str(item["text"]))

            else:
                # Handle objects containing a text attribute
                if hasattr(item, "text"):
                    parts.append(str(item.text))

        text = "\n".join(parts)

    # Final safety conversion
    if not isinstance(text, str):
        text = str(text)

    text = text.strip()

    # Remove accidental markdown code fences
    text = re.sub(
        r"^```(?:markdown|md)?\s*",
        "",
        text,
        flags=re.IGNORECASE
    )

    text = re.sub(
        r"\s*```$",
        "",
        text
    )

    return text.strip()

def analyze_resume(resume_text, job_description):
    """Analyze resume against job description using Gemini."""

    api_key = None

    try:
        api_key = st.secrets.get("GOOGLE_API_KEY")
    except Exception:
        pass

    if not api_key:
        api_key = os.getenv("GOOGLE_API_KEY")

    if not api_key:
        raise Exception(
            "GOOGLE_API_KEY is not configured. "
            "Add it to Streamlit Secrets or your .env file."
        )

    model = ChatGoogleGenerativeAI(
        model="gemini-3.5-flash",
        google_api_key=api_key,
        temperature=0.2,
    )

    prompt = f"""
You are an expert technical recruiter and resume analyst.

Analyze the candidate's resume against the provided job description.

IMPORTANT:
- Use only information actually present in the resume.
- Do not invent skills, experience, education, projects, certifications or achievements.
- Clearly distinguish matching skills from missing skills.
- Give practical and specific suggestions.
- Keep the response professional and useful for job preparation.

Return the analysis using exactly these sections:

# Resume Match Analysis

## 1. Overall Match
Give an estimated match percentage and briefly explain the basis.

## 2. Matching Skills
List the skills from the resume that match the job description.

## 3. Missing Skills
List important job-description skills that are not clearly present in the resume.

## 4. Relevant Experience & Projects
Identify resume projects or experience relevant to the role.

## 5. Resume Strengths
List the strongest parts of the resume for this job.

## 6. Improvement Suggestions
Give concrete changes that could improve the resume for this role.

## 7. Interview Preparation
List the most important technical topics the candidate should prepare based on the job description.

## 8. Final Summary
Give a short factual summary of the candidate's alignment with the role.

-----------------------------
RESUME
-----------------------------

{resume_text}

-----------------------------
JOB DESCRIPTION
-----------------------------

{job_description}
"""

    response = model.invoke(prompt)

    if hasattr(response, "content"):
        return clean_response(response.content)

    return clean_response(response)

# ============================================================
# HERO
# ============================================================

hero_html = (
    '<div class="hero-card">'
    '<div class="hero-title">📄 AI Resume Analyzer</div>'
    '<div class="hero-subtitle">'
    'Analyze your resume against a job description using '
    'LangChain + Google Gemini.'
    '<br>'
    'Identify matching skills, missing skills and practical '
    'improvement suggestions.'
    '</div>'
    '</div>'
)

st.markdown(hero_html, unsafe_allow_html=True)



# ============================================================
# MAIN SECTION
# ============================================================

st.markdown(
    '<div class="section-title">🧑‍💼 Analyze Your Resume</div>',
    unsafe_allow_html=True,
)


# IMPORTANT:
# Streamlit supports gap values such as "small", "medium", "large".
# Do NOT use invalid values like "1.1rem".

left, right = st.columns([1, 1], gap="medium")


# ============================================================
# RESUME UPLOAD
# ============================================================

with left:

    st.markdown(
        '<div class="custom-label">📄 Resume</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="custom-help">Upload your resume</div>',
        unsafe_allow_html=True,
    )

    uploaded_resume = st.file_uploader(
        "Resume",
        type=["pdf", "docx"],
        label_visibility="collapsed",
        help="Upload your resume in PDF or DOCX format.",
    )

    st.caption("200MB per file • PDF, DOCX")


# ============================================================
# JOB DESCRIPTION
# ============================================================

with right:

    st.markdown(
        '<div class="custom-label">💼 Job Description</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="custom-help">Paste the job description</div>',
        unsafe_allow_html=True,
    )

    job_description = st.text_area(
        "Job Description",
        height=180,
        placeholder=(
            "Example:\n\n"
            "We are looking for a Python Developer with SQL, "
            "FastAPI and Git experience..."
        ),
        label_visibility="collapsed",
    )


# ============================================================
# BUTTON
# ============================================================

st.write("")

analyze_clicked = st.button(
    "🚀 Analyze Resume",
    type="primary",
)


# ============================================================
# ANALYSIS
# ============================================================

if analyze_clicked:

    if uploaded_resume is None:
        st.error("Please upload your resume first.")

    elif not job_description.strip():
        st.error("Please paste the job description.")

    else:

        with st.spinner("Analyzing your resume..."):

            try:

                resume_text = extract_resume_text(
                    uploaded_resume
                )

                if not resume_text:
                    st.error(
                        "No readable text was found in the uploaded resume."
                    )
                    st.stop()

                if isinstance(job_description, list):
                     job_description = " ".join(str(x) for x in job_description)
                result = analyze_resume(resume_text,job_description.strip(),)
                st.session_state["analysis_result"] = result

            except Exception as e:
                st.error(str(e))


# ============================================================
# DISPLAY RESULT
# ============================================================

if "analysis_result" in st.session_state:

    st.markdown("---")

    st.markdown(
        '<div class="section-title">📊 Resume Analysis</div>',
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="result-card">',
        unsafe_allow_html=True,
    )

    st.markdown(
        st.session_state["analysis_result"]
    )

    st.markdown(
        "</div>",
        unsafe_allow_html=True,
    )


# ============================================================
# FOOTER
# ============================================================

st.markdown(
    """
    <div class="footer">
        © AI Resume Analyzer
        <br>
        Built with Python • LangChain • Google Gemini • Streamlit
    </div>
    """,
    unsafe_allow_html=True,
)